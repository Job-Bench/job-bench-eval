"""Extract saved Notebook and PowerPoint text without running artifact code."""

from __future__ import annotations

import csv
from html.parser import HTMLParser
import io
import json
from pathlib import Path
import re


class _HTMLText(HTMLParser):
    """Keep visible HTML text and table boundaries, ignoring executable payloads."""

    _HIDDEN = {"script", "style", "template", "head"}
    _BLOCKS = {
        "address", "article", "aside", "blockquote", "div", "dl", "dt", "dd",
        "fieldset", "figcaption", "figure", "footer", "form", "h1", "h2", "h3",
        "h4", "h5", "h6", "header", "hr", "li", "main", "nav", "ol", "p",
        "pre", "section", "table", "ul",
    }

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.hidden: list[str] = []
        self.row_cells: list[int] = []
        self.cell_depth = 0
        self.pre_depth = 0

    def _newline(self):
        if self.parts and not self.parts[-1].endswith("\n"):
            self.parts.append("\n")

    def handle_starttag(self, tag, attrs):
        if tag in self._HIDDEN:
            self.hidden.append(tag)
        if self.hidden:
            return
        if tag == "tr":
            self._newline()
            self.row_cells.append(0)
        elif tag in {"td", "th"}:
            if self.row_cells:
                if self.row_cells[-1]:
                    self.parts.append("\t")
                self.row_cells[-1] += 1
            self.cell_depth += 1
        elif tag in self._BLOCKS or tag == "br":
            self._newline()
        if tag == "pre":
            self.pre_depth += 1

    def handle_endtag(self, tag):
        if self.hidden:
            if tag == self.hidden[-1]:
                self.hidden.pop()
            return
        if tag in {"td", "th"}:
            self.cell_depth = max(0, self.cell_depth - 1)
            if self.parts:
                self.parts[-1] = self.parts[-1].rstrip(" ")
        elif tag == "tr":
            self._newline()
            if self.row_cells:
                self.row_cells.pop()
        elif tag in self._BLOCKS:
            self._newline()
        if tag == "pre":
            self.pre_depth = max(0, self.pre_depth - 1)

    def handle_data(self, data):
        if self.hidden:
            return
        if self.row_cells and not self.cell_depth and not data.strip():
            return
        if not self.pre_depth:
            data = re.sub(r"\s+", " ", data)
            if not self.parts or self.parts[-1].endswith(("\n", "\t")):
                data = data.lstrip(" ")
        self.parts.append(data)

    def text(self):
        return re.sub(r" *\n *", "\n", "".join(self.parts)).strip(" \n")


def _saved_text(value, context: str, parts: list[str], separator: str = "") -> str:
    """Normalize nbformat string/list fields while retaining valid list fragments."""
    if isinstance(value, str):
        return value
    if isinstance(value, list):
        valid = []
        for index, fragment in enumerate(value, 1):
            if isinstance(fragment, str):
                valid.append(fragment)
            else:
                parts.append(f"[WARNING: {context}, item {index}: expected text; skipped]")
        return separator.join(valid)
    parts.append(f"[WARNING: {context}: expected a string or list of strings; skipped]")
    return ""


def _notebook_output(output, context: str, parts: list[str]):
    if not isinstance(output, dict):
        parts.append(f"[WARNING: {context}: expected an output object; skipped]")
        return

    output_type = output.get("output_type")
    if "text" in output:
        stream = _saved_text(output["text"], f"{context} text", parts)
        if stream:
            parts.append(stream)

    if not isinstance(output_type, str):
        parts.append(f"[WARNING: {context}: expected a text output_type; skipped]")
        return

    if output_type == "error":
        name = _saved_text(output.get("ename", ""), f"{context} error name", parts)
        message = _saved_text(output.get("evalue", ""), f"{context} error message", parts)
        summary = ": ".join(item for item in (name, message) if item)
        traceback = _saved_text(output.get("traceback", ""), f"{context} traceback", parts, "\n")
        if summary and summary not in traceback:
            parts.append(summary)
        if traceback:
            parts.append(traceback)

    if output_type not in {"display_data", "execute_result"}:
        return
    data = output.get("data", {})
    if not isinstance(data, dict):
        parts.append(f"[WARNING: {context}: expected a MIME bundle object; skipped]")
        return

    # Different representations may contain complementary information (for
    # example, HTML tables can contain rows elided from a plain-text preview).
    # Deduplicate only identical rendered representations within this output.
    seen: set[str] = set()
    for mime in ("text/plain", "text/markdown", "text/html", "text/latex", "application/json"):
        if mime not in data:
            continue
        try:
            if mime == "application/json":
                text = json.dumps(data[mime], ensure_ascii=False, indent=2)
            else:
                text = _saved_text(data[mime], f"{context} {mime}", parts)
                if mime == "text/html":
                    parser = _HTMLText()
                    parser.feed(text)
                    parser.close()
                    text = parser.text()
            key = text.strip()
            if key and key not in seen:
                seen.add(key)
                parts.append(f"--- {mime} ---\n{text}")
        except Exception as exc:
            parts.append(f"[WARNING: {context} {mime}: could not extract text: {exc}]")


def read_notebook(path: Path) -> str:
    """Read cell sources and already-saved text outputs; never execute a notebook."""
    try:
        notebook = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(notebook, dict) or not isinstance(notebook.get("cells", []), list):
            raise ValueError("expected a notebook object with a cells list")
    except Exception as exc:
        return f"[ERROR: Failed to read notebook {path.name}: {exc}]"

    parts: list[str] = []
    for index, cell in enumerate(notebook.get("cells", []), 1):
        context = f"cell {index}"
        if not isinstance(cell, dict):
            parts.append(f"[WARNING: {context}: expected a cell object; skipped]")
            continue
        parts.append(f"=== {cell.get('cell_type', 'unknown')} ===")
        parts.append(_saved_text(cell.get("source", ""), f"{context} source", parts))
        outputs = cell.get("outputs", [])
        if not isinstance(outputs, list):
            parts.append(f"[WARNING: {context}: expected an outputs list; skipped]")
            continue
        for output_index, output in enumerate(outputs, 1):
            _notebook_output(output, f"{context}, output {output_index}", parts)
    return "\n".join(parts)


def _presentation_shapes(shapes, parts: list[str], context: str):
    for index, shape in enumerate(shapes, 1):
        shape_context = f"{context}, shape {index}"
        try:
            if shape.has_text_frame and shape.text:
                parts.append(shape.text)
            if shape.has_table:
                parts.append(f"=== Table: {shape.name} ===")
                buffer = io.StringIO()
                writer = csv.writer(buffer, delimiter="\t", lineterminator="\n")
                for row in shape.table.rows:
                    # Merged origins own their text; spanned cells retain empty
                    # columns so the table's grid remains visible to the judge.
                    writer.writerow(["" if cell.is_spanned else cell.text for cell in row.cells])
                parts.append(buffer.getvalue().rstrip("\n"))
            if hasattr(shape, "shapes"):
                _presentation_shapes(shape.shapes, parts, shape_context)
        except Exception as exc:
            parts.append(f"[WARNING: {shape_context}: could not extract shape: {exc}]")


def read_presentation(path: Path) -> str:
    """Read slide text and table grids, including shapes inside nested groups."""
    try:
        from pptx import Presentation

        presentation = Presentation(str(path))
        parts: list[str] = []
        for index, slide in enumerate(presentation.slides, 1):
            parts.append(f"=== Slide {index} ===")
            _presentation_shapes(slide.shapes, parts, f"slide {index}")
        return "\n".join(parts)
    except ImportError:
        return f"[ERROR: python-pptx not available for {path.name}]"
    except Exception as exc:
        return f"[ERROR: Failed to read PowerPoint {path.name}: {exc}]"
