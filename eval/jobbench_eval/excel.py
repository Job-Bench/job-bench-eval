"""Coordinate-preserving spreadsheet evidence, including formula provenance."""
from __future__ import annotations

from datetime import date, datetime, time
import json
from pathlib import Path
import re
from zipfile import ZipFile

from .calc import ExtractionError, recalculate as calculate
from .excel_text import format_sheets

MAX_GRID_CELLS = 2_000_000
MAX_UNCOMPRESSED_BYTES = 128 * 1024 * 1024
RESTRICTED_FORMULA = re.compile(
    r"\b(?:WEBSERVICE|DDE|RTD|CALL|REGISTER\.ID|NOW|TODAY|RAND|RANDBETWEEN|RANDARRAY|CELL|INFO)\s*\(",
    re.IGNORECASE,
)


def _value(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def _json(value) -> str:
    return json.dumps(_value(value), ensure_ascii=False, allow_nan=False)


def _restricted_formula(formula: str) -> bool:
    from openpyxl.formula.tokenizer import Tokenizer, TokenizerError
    try:
        tokens = Tokenizer(formula if formula.startswith("=") else "=" + formula).items
    except (TokenizerError, IndexError):
        return True
    return any(token.type == "FUNC" and token.subtype == "OPEN"
               and RESTRICTED_FORMULA.search(token.value) for token in tokens)


def read_excel(path: Path, *, recalculate: bool = True, diagnostics: list | None = None) -> str:
    if path.suffix.lower() == ".xls":
        return _read_legacy_excel(path, diagnostics)
    import openpyxl

    with ZipFile(path) as archive:
        if sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
            raise ExtractionError("Excel expanded content exceeds the 128 MiB extraction limit")
        external = any(name.startswith("xl/externalLinks/") for name in archive.namelist())
        macros = any(name.lower().endswith("vbaproject.bin") for name in archive.namelist())

    formulas = openpyxl.load_workbook(path, data_only=False, read_only=True, keep_links=False)
    cached = openpyxl.load_workbook(path, data_only=True, read_only=True, keep_links=False)
    sheets = []
    requests = {}
    restricted = []
    named_formulas = []
    formula_count = empty_count = zero_count = 0
    try:
        scopes = [("workbook", formulas.defined_names)]
        scopes.extend((sheet.title, sheet.defined_names) for sheet in formulas.worksheets)
        for scope, names in scopes:
            for name, definition in names.items():
                expression = definition.attr_text or ""
                named_formulas.append(f"{scope}/{name}: {expression}")
                if _restricted_formula(expression):
                    restricted.append(f"defined name {scope}/{name}: external, volatile or unsupported expression")
        for sheet in formulas.worksheets:
            values = cached[sheet.title]
            if (sheet.max_row or 0) * (sheet.max_column or 0) > MAX_GRID_CELLS:
                raise ExtractionError(f"Excel sheet {sheet.title!r} exceeds the cell extraction limit")
            records = []
            for source_row, cached_row in zip(sheet.iter_rows(), values.iter_rows()):
                for cell, value_cell in zip(source_row, cached_row):
                    if cell.value is None:
                        continue
                    record = {"cell": cell.coordinate, "value": _value(value_cell.value)}
                    if cell.data_type == "f":
                        formula = cell.value
                        if not isinstance(formula, str):
                            # Array formulas are objects in openpyxl 3.1. Their
                            # anchor formula and declared spill range are evidence.
                            formula = getattr(formula, "text", None)
                            record["range"] = getattr(cell.value, "ref", None)
                            restricted.append(f"{sheet.title}!{cell.coordinate}: array/data-table formula needs range-aware calculation")
                        if not isinstance(formula, str):
                            restricted.append(f"{sheet.title}!{cell.coordinate}: unsupported formula structure")
                            formula = "[unsupported formula structure]"
                        record["formula"] = formula
                        record["date_format"] = openpyxl.styles.numbers.is_date_format(cell.number_format)
                        formula_count += 1
                        empty_count += value_cell.value is None
                        zero_count += value_cell.value == 0
                        requests.setdefault(sheet.title, []).append(cell.coordinate)
                        if _restricted_formula(formula):
                            restricted.append(f"{sheet.title}!{cell.coordinate}: external or volatile function")
                    elif cell.data_type == "e":
                        record["cell_error"] = True
                    records.append(record)
            sheets.append((sheet.title, records))
    finally:
        formulas.close()
        cached.close()

    warnings = []
    if external:
        warnings.append("Workbook contains external links; they are not refreshed.")
    if macros:
        warnings.append("Workbook contains macros; they are not executed.")
    if restricted:
        warnings.append("External/volatile or unsupported formulas prevent deterministic recalculation: "
                        + "; ".join(restricted[:20]))

    result = None
    status = "no_formulas"
    if formula_count:
        if external or macros or restricted:
            status = "not_recalculated_restricted"
        elif not recalculate:
            status = "not_recalculated"
        else:
            result = calculate(path, requests)
            status = "recalculated"
        if result is None:
            warnings.append("Formula values below are saved caches, not independently calculated results. "
                            "A null cache does not mean the formula is absent; zero may be a placeholder. "
                            "Numeric correctness is unverified.")

    calculation_error_count = sum(bool(cell["error_code"])
                                  for cells in result["cells"].values() for cell in cells.values()) if result else 0
    if calculation_error_count:
        warnings.append(f"{calculation_error_count} formula cells returned LibreOffice calculation errors. "
                        "These can reflect workbook/formula defects or engine compatibility differences. "
                        "Numeric correctness of those cells remains unverified; retain the original "
                        "formulas and caches as evidence rather than treating an engine error alone as proof of a wrong answer.")

    parts = [f"=== Excel calculation: {status} ==="]
    if result:
        parts.append(f"Engine: {result['engine']} ({result['version']}); network disabled.")
    parts.extend("[WARNING: " + warning + "]" for warning in warnings)
    if named_formulas:
        parts.append("=== Defined names ===\n" + "\n".join(named_formulas))
    evidence = []
    for sheet_name, records in sheets:
        for record in records:
            if "formula" in record:
                item = {"sheet": sheet_name, **record}
                if result:
                    item["recalculated"] = result["cells"][sheet_name][record["cell"]]
                evidence.append(item)
    header = "\n".join(parts)
    if len(header) > 9000:
        header = header[:4500] + "\n[EXCEL TEXT TRUNCATED: workbook metadata; full metadata in diagnostics]\n" + header[-4000:]
    text = header + "\n" + format_sheets(sheets, result)
    if diagnostics is not None:
        diagnostics.append({"type": "excel", "status": status, "formula_count": formula_count,
                            "calculation_error_count": calculation_error_count,
                            "empty_cached_values": empty_count, "zero_cached_values": zero_count,
                            "engine": result["engine"] if result else None, "warnings": warnings,
                            "defined_names": named_formulas, "formulas": evidence,
                            "text_truncated": "[EXCEL TEXT TRUNCATED" in text,
                            "truncation_notes": [line for line in text.splitlines()
                                                 if "[EXCEL TEXT TRUNCATED" in line]})
    return text


def _read_legacy_excel(path: Path, diagnostics: list | None) -> str:
    import xlrd
    from openpyxl.utils import get_column_letter
    workbook = xlrd.open_workbook(path, on_demand=True)
    warning = "Legacy XLS: saved values only; formula expressions and recalculation are unavailable."
    parts = ["[WARNING: " + warning + "]"]
    try:
        for sheet in workbook.sheets():
            if sheet.nrows * sheet.ncols > MAX_GRID_CELLS:
                raise ExtractionError("Legacy Excel sheet exceeds the cell extraction limit")
            parts.append(f"=== Sheet: {sheet.name} ===")
            for row in range(sheet.nrows):
                for col in range(sheet.ncols):
                    cell = sheet.cell(row, col)
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        continue
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, workbook.datemode)
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        value = xlrd.error_text_from_code.get(value, "[unknown Excel error]")
                    parts.append(f"{get_column_letter(col + 1)}{row + 1}: {_json(value)}")
    finally:
        workbook.release_resources()
    if diagnostics is not None:
        diagnostics.append({"type": "excel", "status": "legacy_cached_values", "warnings": [warning]})
    return "\n".join(parts)
