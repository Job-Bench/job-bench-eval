from __future__ import annotations

import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

import openpyxl

JUDGE = Path(__file__).resolve().parents[1] / "eval" / "judge.py"


class JudgeEvidenceIntegrationTests(unittest.TestCase):
    def write_current_partial_details(self, details):
        implementation = json.loads(subprocess.check_output(
            [sys.executable, str(JUDGE), "--print-implementation"], text=True,
        ))
        original = json.dumps({
            "judge_implementation": implementation,
            "evaluated_model": "saved-model", "judge_model": "offline-judge",
            "total_score": 1, "max_score": 2, "passed_count": 1, "total_count": 2,
            "rubrics": [{"index": 0, "weight": 1, "result": {"passed": True, "score": 1}}],
        }, indent=2) + "\n"
        details.write_text(original)
        return original

    def test_resume_extraction_failure_preserves_same_implementation_partial_scores(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            output = root / "output"
            output.mkdir()
            book = openpyxl.Workbook()
            book.active["A1"] = "=1+1"
            book.save(output / "book.xlsx")
            rubrics = root / "rubrics.json"
            rubrics.write_text(json.dumps({"rubrics": [
                {"rubric": "First criterion", "weight": 1},
                {"rubric": "Second criterion", "weight": 1},
            ]}))
            details = root / "partial.json"
            original = self.write_current_partial_details(details)
            reward = root / "reward.json"
            report = root / "extraction.json"

            result = subprocess.run([sys.executable, str(JUDGE), "--output-dir", str(output),
                "--rubrics-file", str(rubrics), "--details-file", str(details),
                "--result-file", str(reward), "--extraction-file", str(report)],
                env={**os.environ, "PATH": "/nonexistent", "JOBBENCH_CALC_NATIVE": "0"},
                capture_output=True, text=True)

            self.assertEqual(result.returncode, 2, result.stderr)
            self.assertEqual(details.read_text(), original)
            self.assertFalse(reward.exists())
            diagnostics = json.loads(report.read_text())
            self.assertEqual(diagnostics["status"], "error")
            self.assertIn("book.xlsx", diagnostics["error"])

    def test_resume_generic_failure_preserves_partial_scores_without_zero_reward(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            output = root / "output"
            output.mkdir()
            rubrics = root / "rubrics.json"
            rubrics.write_text('{"rubrics":[')
            details = root / "partial.json"
            original = self.write_current_partial_details(details)
            reward = root / "reward.json"
            report = root / "extraction.json"

            result = subprocess.run([sys.executable, str(JUDGE), "--output-dir", str(output),
                "--rubrics-file", str(rubrics), "--details-file", str(details),
                "--result-file", str(reward), "--extraction-file", str(report)],
                capture_output=True, text=True)

            self.assertEqual(result.returncode, 2, result.stderr)
            self.assertEqual(details.read_text(), original)
            self.assertFalse(reward.exists())
            diagnostics = json.loads(report.read_text())
            self.assertEqual(diagnostics["status"], "error")
            self.assertIn("error", diagnostics)

    def test_resume_missing_or_empty_rubrics_preserves_partial_scores(self):
        for rubric_content in (None, '{"rubrics": []}'):
            with self.subTest(rubric_content=rubric_content), tempfile.TemporaryDirectory() as temporary:
                root = Path(temporary)
                output = root / "output"
                output.mkdir()
                rubrics = root / "rubrics.json"
                if rubric_content is not None:
                    rubrics.write_text(rubric_content)
                details = root / "partial.json"
                original = self.write_current_partial_details(details)
                reward = root / "reward.json"
                report = root / "extraction.json"

                result = subprocess.run([sys.executable, str(JUDGE), "--output-dir", str(output),
                    "--rubrics-file", str(rubrics), "--details-file", str(details),
                    "--result-file", str(reward), "--extraction-file", str(report)],
                    capture_output=True, text=True)

                self.assertEqual(result.returncode, 2, result.stderr)
                self.assertEqual(details.read_text(), original)
                self.assertFalse(reward.exists())
                diagnostics = json.loads(report.read_text())
                self.assertEqual(diagnostics["status"], "error")
                self.assertIn("rubrics", diagnostics["error"].lower())

    def test_malformed_historical_details_are_not_overwritten(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            output = root / "output"
            output.mkdir()
            rubric = root / "rubrics.json"
            rubric.write_text('{"rubrics":[{"rubric":"correct", "weight":1}]}')
            details = root / "old.json"
            original = '{"rubrics":['
            details.write_text(original)
            result = subprocess.run([sys.executable, str(JUDGE), "--output-dir", str(output),
                "--rubrics-file", str(rubric), "--details-file", str(details)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 2)
            self.assertEqual(details.read_text(), original)

    def test_regular_file_is_not_an_empty_output_directory(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            output = root / "output"
            output.write_text("not a directory")
            report = root / "extraction.json"
            result = subprocess.run([sys.executable, str(JUDGE), "--output-dir", str(output),
                "--extract-only", "--extraction-file", str(report)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 2)
            self.assertFalse(report.exists())

    def test_existing_reward_is_not_mixed_with_new_details(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            output = root / "output"
            output.mkdir()
            rubric = root / "rubrics.json"
            rubric.write_text('{"rubrics":[{"rubric":"correct", "weight":1}]}')
            reward = root / "reward.json"
            reward.write_text('{"reward":0.9}')
            details = root / "new.json"
            result = subprocess.run([sys.executable, str(JUDGE), "--output-dir", str(output),
                "--rubrics-file", str(rubric), "--result-file", str(reward), "--details-file", str(details)],
                capture_output=True, text=True)
            self.assertEqual(result.returncode, 2)
            self.assertEqual(reward.read_text(), '{"reward":0.9}')
            self.assertFalse(details.exists())

    def test_extract_only_failure_never_changes_scoring_destinations(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            output = root / "output"
            output.mkdir()
            book = openpyxl.Workbook()
            book.active["A1"] = "=1+1"
            book.save(output / "book.xlsx")
            details = root / "old.json"
            original = '{"rubrics":[{"index":0,"result":{"passed":true}}]}'
            details.write_text(original)
            result = subprocess.run([sys.executable, str(JUDGE), "--output-dir", str(output),
                "--extract-only", "--details-file", str(details), "--extraction-file", str(root / "extraction.json")],
                env={**os.environ, "PATH": "/nonexistent", "JOBBENCH_CALC_NATIVE": "0"}, capture_output=True, text=True)
            self.assertEqual(result.returncode, 2)
            self.assertEqual(details.read_text(), original)

    def test_older_partial_judge_results_are_preserved_without_api_calls(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            output = root / "output"
            output.mkdir()
            (output / "answer.txt").write_text("answer")
            rubric = root / "rubrics.json"
            rubric.write_text('{"rubrics":[{"rubric":"correct", "weight":1}]}')
            details = root / "old.json"
            original = '{"rubrics":[{"index":0,"result":{"passed":true}}]}'
            details.write_text(original)
            result = subprocess.run([sys.executable, str(JUDGE), "--output-dir", str(output),
                "--rubrics-file", str(rubric), "--details-file", str(details)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 2)
            self.assertIn("JUDGE_RUN_LABEL", result.stderr)
            self.assertEqual(details.read_text(), original)
            self.assertFalse(details.with_suffix(".extraction.json").exists())

    def test_extract_only_records_saved_notebook_results_and_provenance(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            output = root / "output"
            output.mkdir()
            notebook = output / "saved.ipynb"
            notebook.write_text(json.dumps({"cells": [{"cell_type": "code", "source": "df",
                "outputs": [{"output_type": "execute_result", "data": {"text/plain": "FINAL_REVENUE_9876"}}]}]}))
            report = root / "extraction.json"
            result = subprocess.run([sys.executable, str(JUDGE), "--output-dir", str(output),
                "--extract-only", "--extraction-file", str(report)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            evidence = json.loads(report.read_text())
            self.assertIn("FINAL_REVENUE_9876", evidence["text"])
            self.assertEqual(evidence["files"][0]["path"], "saved.ipynb")
            self.assertEqual(len(evidence["implementation"]["sha256"]), 64)
            self.assertEqual(evidence["status"], "complete")

    def test_engine_failure_produces_diagnostics_and_no_zero_reward(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            output = root / "output"
            output.mkdir()
            book = openpyxl.Workbook()
            book.active["A1"] = "=1+1"
            book.save(output / "book.xlsx")
            rubrics = root / "rubrics.json"
            rubrics.write_text(json.dumps({"rubrics": [{"rubric": "Answer is two", "weight": 1}]}))
            reward = root / "reward.json"
            report = root / "extraction.json"
            details = root / "details.json"
            result = subprocess.run([sys.executable, str(JUDGE), "--output-dir", str(output),
                "--rubrics-file", str(rubrics), "--result-file", str(reward), "--details-file", str(details),
                "--extraction-file", str(report)], env={**os.environ, "PATH": "/nonexistent", "JOBBENCH_CALC_NATIVE": "0"},
                capture_output=True, text=True)
            self.assertEqual(result.returncode, 2, result.stderr)
            self.assertFalse(reward.exists())
            self.assertEqual(json.loads(report.read_text())["status"], "error")
            self.assertEqual(json.loads(details.read_text())["error_type"], "extraction_error")


if __name__ == "__main__":
    unittest.main()
