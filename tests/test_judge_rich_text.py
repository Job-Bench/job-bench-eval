"""Text extraction regressions using saved Notebook and PowerPoint artifacts."""

import json
from pathlib import Path
import sys
import tempfile
import unittest

from pptx import Presentation
from pptx.util import Inches


sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "eval"))
from jobbench_eval.rich_text import read_notebook, read_presentation


class RichTextArtifactTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)

    def notebook(self, cells):
        artifact = self.root / "results.ipynb"
        artifact.write_text(json.dumps({
            "nbformat": 4, "nbformat_minor": 5, "metadata": {}, "cells": cells,
        }), encoding="utf-8")
        return artifact

    def code_cell(self, outputs, source="result"):
        return {
            "cell_type": "code", "metadata": {}, "source": source,
            "execution_count": 1, "outputs": outputs,
        }

    def test_notebook_preserves_sources_streams_without_execution(self):
        marker = self.root / "must-not-execute"
        source = f"from pathlib import Path\nPath({str(marker)!r}).touch()"
        artifact = self.notebook([
            {"cell_type": "markdown", "metadata": {}, "source": ["# Report\n", "Saved analysis"]},
            {"cell_type": "raw", "metadata": {}, "source": "Raw observation"},
            self.code_cell([
                {"output_type": "stream", "name": "stdout", "text": ["First line\n", "Second line\n"]},
                {"output_type": "stream", "name": "stderr", "text": "Saved warning\n"},
            ], source=source),
        ])

        text = read_notebook(artifact)

        for expected in ("=== markdown ===", "# Report\nSaved analysis", "Raw observation",
                         source, "First line\nSecond line", "Saved warning"):
            self.assertIn(expected, text)
        self.assertFalse(marker.exists())

    def test_notebook_keeps_html_table_details_beyond_plain_text_preview(self):
        artifact = self.notebook([self.code_cell([{
            "output_type": "execute_result", "execution_count": 1, "metadata": {},
            "data": {
                "text/plain": ["  Company ... Revenue\n", "[2 rows x 2 columns]"],
                "text/html": [
                    "<style>STYLE_SENTINEL</style><script>SCRIPT_SENTINEL</script>",
                    "<table><tr><th>Company</th><th>Revenue</th></tr>",
                    "<tr><td>Alpha &amp; Beta</td><td>451.2</td></tr>",
                    "<tr><td>Gamma</td><td>918.4</td></tr></table>",
                    '<img src="data:image/png;base64,BASE64_SENTINEL">',
                ],
                "image/png": "IMAGE_PAYLOAD_SENTINEL",
                "application/javascript": "JAVASCRIPT_SENTINEL",
            },
        }])])

        text = read_notebook(artifact)

        self.assertIn("[2 rows x 2 columns]", text)
        self.assertIn("Company\tRevenue", text)
        self.assertIn("Alpha & Beta\t451.2", text)
        self.assertIn("Gamma\t918.4", text)
        for excluded in ("STYLE_SENTINEL", "SCRIPT_SENTINEL", "BASE64_SENTINEL",
                         "IMAGE_PAYLOAD_SENTINEL", "JAVASCRIPT_SENTINEL", "<table>"):
            self.assertNotIn(excluded, text)

    def test_notebook_keeps_all_supported_saved_rich_mime_representations(self):
        artifact = self.notebook([self.code_cell([
            {"output_type": "display_data", "metadata": {}, "data": {
                "text/markdown": ["## Finding\n", "**Profit increased**"],
                "text/latex": ["\\frac{17}", "{23}"],
                "application/json": {"count": 17, "labels": ["café", "東京"]},
            }},
            {"output_type": "execute_result", "execution_count": 1, "metadata": {},
             "data": {"text/plain": "Result scalar: 73"}},
        ])])

        text = read_notebook(artifact)

        for expected in ("## Finding\n**Profit increased**", "\\frac{17}{23}",
                         '"count": 17', '"labels":', "café", "東京", "Result scalar: 73"):
            self.assertIn(expected, text)

    def test_notebook_deduplicates_equivalent_mime_text_within_one_output(self):
        artifact = self.notebook([self.code_cell([
            {"output_type": "display_data", "metadata": {}, "data": {
                "text/plain": "Unique finding 42", "text/markdown": ["Unique finding 42"],
                "text/html": "<p>Unique finding 42</p>",
            }},
            {"output_type": "stream", "name": "stdout", "text": "repeated log\n"},
            {"output_type": "stream", "name": "stdout", "text": "repeated log\n"},
        ])])

        text = read_notebook(artifact)

        self.assertEqual(text.count("Unique finding 42"), 1)
        self.assertEqual(text.count("repeated log"), 2)

    def test_notebook_preserves_saved_error_name_message_and_traceback(self):
        artifact = self.notebook([self.code_cell([{
            "output_type": "error", "ename": "ValueError", "evalue": "invalid sample",
            "traceback": ["Traceback (most recent call last):", "Cell In[1], line 7", "ValueError: invalid sample"],
        }])])

        text = read_notebook(artifact)

        self.assertIn("ValueError: invalid sample", text)
        self.assertIn("Traceback (most recent call last):\nCell In[1], line 7", text)

    def test_notebook_malformed_outputs_warn_without_hiding_remaining_content(self):
        artifact = self.notebook([
            self.code_cell([
                None,
                {"output_type": "stream", "name": "stdout", "text": 123},
                {"output_type": "display_data", "metadata": {}, "data": ["invalid bundle"]},
                {"output_type": "display_data", "metadata": {}, "data": {
                    "text/plain": {"invalid": "value"},
                    "text/html": "<p>Surviving HTML result</p>",
                }},
                {"output_type": "stream", "name": "stdout", "text": "Surviving stream"},
            ], source="First valid source"),
            self.code_cell([], source=["Later valid source"]),
        ])

        text = read_notebook(artifact)

        for expected in ("First valid source", "Surviving HTML result", "Surviving stream", "Later valid source"):
            self.assertIn(expected, text)
        self.assertIn("WARNING", text)
        self.assertNotIn("[ERROR: Failed to read notebook", text)

    def test_notebook_corrupt_file_returns_readable_error(self):
        artifact = self.root / "broken.ipynb"
        artifact.write_text("not valid JSON", encoding="utf-8")

        text = read_notebook(artifact)

        self.assertIn("ERROR", text)
        self.assertIn("broken.ipynb", text)

    def test_notebook_bad_output_type_and_cell_fields_do_not_abort_later_cells(self):
        artifact = self.notebook([
            None,
            self.code_cell([
                {"output_type": ["not", "a", "string"], "data": {"text/plain": "invalid output"}},
                {"output_type": "stream", "name": "stdout", "text": ["Valid fragment ", 8, "survives"]},
            ], source=["Valid source ", None, "survives"]),
            self.code_cell({"not": "an output list"}, source="Source with malformed outputs"),
            self.code_cell([], source="Final intact cell"),
        ])

        text = read_notebook(artifact)

        for expected in ("Valid source survives", "Valid fragment survives",
                         "Source with malformed outputs", "Final intact cell"):
            self.assertIn(expected, text)
        self.assertIn("WARNING", text)
        self.assertIn("cell 2", text)

    def test_presentation_extracts_table_grid_with_merged_origins_once(self):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text = "Quarterly results"
        table = slide.shapes.add_table(3, 3, 0, Inches(1), Inches(6), Inches(3)).table
        table.cell(0, 0).merge(table.cell(0, 1))
        table.cell(0, 0).text = "Merged performance heading"
        table.cell(0, 2).text = "Units"
        for column, value in enumerate(("Revenue", "123.5", "USD")):
            table.cell(1, column).text = value
        for column, value in enumerate(("Profit", "27.8", "USD")):
            table.cell(2, column).text = value
        artifact = self.root / "table.pptx"
        presentation.save(artifact)

        text = read_presentation(artifact)

        self.assertIn("=== Slide 1 ===", text)
        self.assertIn("Quarterly results", text)
        self.assertEqual(text.count("Merged performance heading"), 1)
        self.assertIn("Merged performance heading\t\tUnits", text)
        self.assertIn("Revenue\t123.5\tUSD", text)
        self.assertIn("Profit\t27.8\tUSD", text)

    def test_presentation_extracts_nested_group_text_and_tables_in_slide_order(self):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        outer = slide.shapes.add_group_shape()
        outer.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text = "Grouped title"
        nested = outer.shapes.add_group_shape()
        nested.shapes.add_textbox(0, Inches(1), Inches(4), Inches(1)).text = "Nested conclusion"
        frame = slide.shapes.add_table(1, 2, 0, Inches(2), Inches(4), Inches(1))
        frame.table.cell(0, 0).text = "Nested metric"
        frame.table.cell(0, 1).text = "894"
        # OOXML supports tables in groups even though GroupShapes has no add_table().
        nested.shapes._spTree.insert_element_before(frame._element, "p:extLst")
        second = presentation.slides.add_slide(presentation.slide_layouts[6])
        second.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text = "Second slide"
        artifact = self.root / "groups.pptx"
        presentation.save(artifact)

        text = read_presentation(artifact)

        for expected in ("Grouped title", "Nested conclusion", "Nested metric\t894", "=== Slide 2 ===", "Second slide"):
            self.assertIn(expected, text)
        self.assertLess(text.index("Nested conclusion"), text.index("=== Slide 2 ==="))
        self.assertEqual(text.count("Nested metric"), 1)

    def test_presentation_corrupt_file_returns_readable_error(self):
        artifact = self.root / "broken.pptx"
        artifact.write_text("not a zip", encoding="utf-8")

        text = read_presentation(artifact)

        self.assertIn("ERROR", text)
        self.assertIn("broken.pptx", text)


if __name__ == "__main__":
    unittest.main()
