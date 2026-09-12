from __future__ import annotations

import sys
import os
from datetime import datetime
from unittest.mock import patch
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import tempfile
import unittest
from pathlib import Path

import openpyxl

EVAL = Path(__file__).resolve().parents[1] / "eval"
sys.path.insert(0, str(EVAL))
from jobbench_eval.excel import read_excel
from jobbench_eval.calc import ExtractionError



class ExcelEvidenceTests(unittest.TestCase):
    def test_formula_only_rows_and_missing_cache_remain_visible(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "book.xlsx"
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = "Results"
            sheet.append(["Input", "Computed"])
            sheet.append([10, "=A2*2"])
            sheet["B3"] = "=SUM(A2:A2)"
            book.save(path)
            before = path.read_bytes()
            text = read_excel(path, recalculate=False)
            self.assertIn("=A2*2", text)
            self.assertIn("B3", text)
            self.assertIn("=SUM(A2:A2)", text)
            self.assertEqual(path.read_bytes(), before)

    def test_literals_dates_errors_and_positions_need_no_calculator(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "literals.xlsx"
            book = openpyxl.Workbook()
            sheet = book.active
            sheet["A1"] = "Repeated heading"
            sheet["B1"] = "Repeated heading"
            sheet["C4"] = 0
            sheet["D4"] = False
            sheet["E4"] = datetime(2025, 1, 2)
            sheet["F4"] = "#DIV/0!"
            book.save(path)
            with patch("jobbench_eval.excel.calculate", side_effect=AssertionError("unexpected calculator")):
                text = read_excel(path)
            for value in ('row,A,B,C,D,E,F', '1,Repeated heading,Repeated heading',
                          '4,,,0,False,2025-01-02T00:00:00,#DIV/0!'):
                self.assertIn(value, text)

    def test_volatile_formula_is_explicitly_unverified_not_fabricated(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "volatile.xlsx"
            book = openpyxl.Workbook()
            book.active["A1"] = "=TODAY()"
            book.save(path)
            records = []
            with patch("jobbench_eval.excel.calculate", side_effect=AssertionError("unexpected calculator")):
                text = read_excel(path, diagnostics=records)
            self.assertIn("=TODAY()", text)
            self.assertIn("cached=null", text)
            self.assertIn("Numeric correctness is unverified", text)
            self.assertEqual(records[0]["status"], "not_recalculated_restricted")

    def test_formula_error_is_not_replaced_by_zero(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "error.xlsx"
            book = openpyxl.Workbook()
            book.active["A1"] = "=1/0"
            book.save(path)
            result = {"engine": "fixture", "version": "fixture", "cells": {
                "Sheet": {"A1": {"value": "#DIV/0!", "display": "#DIV/0!", "error_code": 532}}}}
            diagnostics = []
            with patch("jobbench_eval.excel.calculate", return_value=result):
                text = read_excel(path, diagnostics=diagnostics)
            self.assertIn("Calc error 532", text)
            self.assertIn('cached=null', text)
            self.assertIn('#DIV/0!', text)
            self.assertIn('compatibility differences', text)
            self.assertIn('Numeric correctness of those cells remains unverified', text)
            self.assertEqual(diagnostics[0]['calculation_error_count'], 1)

    def test_unavailable_engine_is_an_extraction_error(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "book.xlsx"
            book = openpyxl.Workbook()
            book.active["A1"] = "=1+1"
            book.save(path)
            with patch("jobbench_eval.excel.calculate", side_effect=ExtractionError("unavailable")):
                with self.assertRaisesRegex(ExtractionError, "unavailable"):
                    read_excel(path)

    def test_named_volatile_formula_is_not_recalculated(self):
        from openpyxl.workbook.defined_name import DefinedName
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "named.xlsx"
            book = openpyxl.Workbook()
            book.defined_names.add(DefinedName("Clock", attr_text="NOW()"))
            book.active["A1"] = "=Clock"
            book.save(path)
            with patch("jobbench_eval.excel.calculate", side_effect=AssertionError("volatile name executed")):
                text = read_excel(path)
            self.assertIn("not_recalculated_restricted", text)

    def test_malformed_formula_keeps_other_cells_and_explains_restriction(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "malformed.xlsx"
            book = openpyxl.Workbook()
            book.active["A1"] = "=1)"
            book.active["B2"] = "other evidence survives"
            book.save(path)
            text = read_excel(path)
            self.assertIn("=1)", text)
            self.assertIn("other evidence survives", text)
            self.assertIn("not_recalculated_restricted", text)

    def test_array_formula_does_not_claim_only_anchor_is_complete_calculation(self):
        from openpyxl.worksheet.formula import ArrayFormula
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "array.xlsx"
            book = openpyxl.Workbook()
            book.active["A1"] = ArrayFormula("A1:A3", "=ROW(A1:A3)")
            book.save(path)
            with patch("jobbench_eval.excel.calculate", side_effect=AssertionError("incomplete array calculation")):
                text = read_excel(path)
            self.assertIn("A1:A3", text)
            self.assertIn("not_recalculated_restricted", text)


@unittest.skipUnless(os.environ.get("JOBBENCH_TEST_CALC") == "1", "requires prepared Docker calculation runtime")
class ExcelRuntimeTests(unittest.TestCase):
    def test_missing_zero_and_stale_caches_are_explicitly_recalculated(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "book.xlsx"
            book = openpyxl.Workbook()
            inputs = book.active
            inputs.title = "Inputs"
            inputs["A1"] = 10
            results = book.create_sheet("Results")
            results["A1"] = "=Inputs!A1*2"
            results["A2"] = "=Inputs!A1*3"
            results["A3"] = "=Inputs!A1*4"
            results["A4"] = '=IF(Inputs!A1=10,"yes","no")'
            book.save(path)
            # Independent OOXML cache fixtures: missing, placeholder zero, stale 999.
            with ZipFile(path) as archive:
                members = {name: archive.read(name) for name in archive.namelist()}
            ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
            tree = ET.fromstring(members["xl/worksheets/sheet2.xml"])
            for address, value in [("A2", "0"), ("A3", "999")]:
                cell = tree.find(f'.//s:c[@r="{address}"]', ns)
                cell.find("s:v", ns).text = value
            members["xl/worksheets/sheet2.xml"] = ET.tostring(tree)
            with ZipFile(path, "w") as archive:
                for name, data in members.items():
                    archive.writestr(name, data)
            before = path.read_bytes()
            records = []
            text = read_excel(path, diagnostics=records)
            evidence = {cell["cell"]: cell for cell in records[0]["formulas"]}
            self.assertEqual([evidence[a]["value"] for a in ("A1", "A2", "A3")], [None, 0, 999])
            self.assertEqual([evidence[a]["recalculated"]["value"] for a in ("A1", "A2", "A3", "A4")],
                             [20, 30, 40, "yes"])
            self.assertIn("network disabled", text)
            self.assertEqual(path.read_bytes(), before)

    def test_timeout_is_reported_without_a_score_or_input_change(self):
        from jobbench_eval.calc import recalculate
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "book.xlsx"
            book = openpyxl.Workbook()
            book.active["A1"] = "=1+1"
            book.save(path)
            before = path.read_bytes()
            with self.assertRaisesRegex(ExtractionError, "exceeded"):
                recalculate(path, {"Sheet": ["A1"]}, timeout=0.001)
            self.assertEqual(path.read_bytes(), before)


if __name__ == "__main__":
    unittest.main()
