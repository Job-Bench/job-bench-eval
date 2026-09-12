"""Deterministic OpenAI-compatible endpoint for local OpenCode Harbor tests.

The endpoint accepts one synthetic model only. It never forwards requests and
does not read or record authorization headers.
"""

from __future__ import annotations

import argparse
import json
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path


DEFAULT_REQUEST_MODEL = "fixture-wire-model"

_INSPECTION_SCRIPT = r"""
import csv
import json
import sqlite3
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import pdfplumber
from openpyxl import load_workbook
from PIL import Image

root = Path(__TASK_ROOT__)
instructions = root / 'TASK_INSTRUCTIONS.txt'
text = instructions.read_text(encoding='utf-8')
checks = []
for path in sorted(item for item in root.rglob('*') if item.is_file()):
    suffix = path.suffix.lower()
    status = 'listed'
    counts = {}
    if suffix == '.csv':
        with path.open(encoding='utf-8-sig', errors='replace', newline='') as stream:
            first_row = next(csv.reader(stream), [])
        counts = {'first_row_cells': len(first_row), 'rows_sampled': int(bool(first_row))}
        status = 'csv-parsed'
    elif suffix == '.xlsx':
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.worksheets[0]
            first_row = next(sheet.iter_rows(max_row=1, values_only=True), ())
            counts = {
                'sheets': len(workbook.worksheets),
                'first_sheet_rows': sheet.max_row,
                'first_row_cells': len(first_row),
            }
        finally:
            workbook.close()
        status = 'xlsx-parsed'
    elif suffix == '.pdf':
        with pdfplumber.open(path) as document:
            first_page_text = document.pages[0].extract_text() or ''
            counts = {'pages': len(document.pages), 'first_page_chars': len(first_page_text)}
        status = 'pdf-parsed'
    elif suffix == '.docx':
        with zipfile.ZipFile(path) as archive:
            document_xml = archive.read('word/document.xml')
        document = ElementTree.fromstring(document_xml)
        namespace = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
        counts = {'paragraphs': len(document.findall(f'.//{{{namespace}}}p'))}
        status = 'docx-parsed'
    elif suffix in {'.png', '.jpg', '.jpeg', '.tif', '.tiff'}:
        with Image.open(path) as image:
            width, height = image.size
            image.load()
        counts = {'width': width, 'height': height}
        status = 'image-parsed'
    elif suffix in {'.db', '.sqlite', '.sqlite3'}:
        uri = f"file:{path}?mode=ro"
        with sqlite3.connect(uri, uri=True) as connection:
            connection.execute('PRAGMA query_only = ON')
            tables = connection.execute(
                "SELECT COUNT(*) FROM sqlite_master WHERE type = 'table'"
            ).fetchone()[0]
        counts = {'tables': tables}
        status = 'sqlite-parsed'
    elif suffix == '.stl':
        if not path.read_bytes()[:80]:
            raise ValueError(f'empty stl: {path}')
        counts = {'header_bytes_sampled': min(path.stat().st_size, 80)}
        status = 'stl-parsed'
    checks.append({'path': str(path.relative_to(root)), 'status': status, 'counts': counts})

output = Path(__OUTPUT_ROOT__)
output.mkdir(parents=True, exist_ok=True)
(output / 'smoke.txt').write_text(
    'Deterministic local OpenCode fixture. Not a benchmark solution.\n',
    encoding='utf-8',
)
(output / 'smoke.json').write_text(
    json.dumps({'instructions_chars': len(text), 'files': checks}, indent=2) + '\n',
    encoding='utf-8',
)
print(json.dumps({'inspected_files': len(checks)}))
"""


def build_inspection_command(
    task_root: str | Path = "/workspace/task_folder",
    output_root: str | Path = "/workspace/output",
) -> str:
    script = _INSPECTION_SCRIPT.replace("__TASK_ROOT__", json.dumps(str(task_root)))
    script = script.replace("__OUTPUT_ROOT__", json.dumps(str(output_root)))
    return "python3 - <<'PY'\n" + script.strip() + "\nPY"


_INSPECTION_COMMAND = build_inspection_command()


def _chunk(delta: dict, *, finish_reason: str | None = None) -> dict:
    return {
        "id": "chatcmpl-jobbench-local",
        "object": "chat.completion.chunk",
        "created": 0,
        "model": DEFAULT_REQUEST_MODEL,
        "choices": [{"index": 0, "delta": delta, "finish_reason": finish_reason}],
    }


def _sse(events: list[dict]) -> bytes:
    lines = [f"data: {json.dumps(event, separators=(',', ':'))}" for event in events]
    lines.append("data: [DONE]")
    return ("\n\n".join(lines) + "\n\n").encode()


class LocalOpenCodeHandler(BaseHTTPRequestHandler):
    server_version = "JobBenchLocalOpenCodeFixture/1"

    def log_message(self, format: str, *args) -> None:
        return

    def _write(self, status: int, content_type: str, body: bytes) -> None:
        self.send_response(status)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def _write_json(self, status: int, value: dict) -> None:
        self._write(status, "application/json", json.dumps(value).encode())

    def do_GET(self) -> None:
        if self.path != "/health":
            self.send_error(404)
            return
        self._write_json(200, {"ok": True})

    def do_POST(self) -> None:
        if self.path != "/v1/chat/completions":
            self.send_error(404)
            return
        try:
            length = int(self.headers.get("Content-Length", "0"))
            payload = json.loads(self.rfile.read(length))
        except (TypeError, ValueError, json.JSONDecodeError):
            self._write_json(400, {"error": {"message": "invalid JSON"}})
            return

        if payload.get("model") != self.server.expected_model:
            self._write_json(
                422,
                {
                    "error": {
                        "message": "expected modelAliases to rewrite fixture-model to fixture-wire-model"
                    }
                },
            )
            return
        if self.server.mode == "provider-error":
            self._write_json(401, {"error": {"message": "synthetic fixture provider error"}})
            return

        self.server.requests.append(
            {
                "model": payload["model"],
                "stream": bool(payload.get("stream")),
                "roles": [item.get("role") for item in payload.get("messages", [])],
            }
        )

        if not payload.get("stream"):
            body = {
                "id": "chatcmpl-jobbench-title",
                "object": "chat.completion",
                "created": 0,
                "model": payload["model"],
                "choices": [
                    {
                        "index": 0,
                        "message": {
                            "role": "assistant",
                            "content": "JobBench local OpenCode fixture",
                        },
                        "finish_reason": "stop",
                    }
                ],
                "usage": {"prompt_tokens": 0, "completion_tokens": 0, "total_tokens": 0},
            }
            self._write_json(200, body)
            return

        roles = [message.get("role") for message in payload.get("messages", [])]
        if "tool" not in roles:
            tool_call = {
                "index": 0,
                "id": "call_jobbench_local_smoke",
                "type": "function",
                "function": {
                    "name": "bash",
                    "arguments": json.dumps(
                        {
                            "command": _INSPECTION_COMMAND,
                            "description": "Inspect fixture inputs and write deterministic smoke artifacts",
                        },
                        separators=(",", ":"),
                    ),
                },
            }
            events = [
                _chunk({"role": "assistant"}),
                _chunk({"tool_calls": [tool_call]}),
                _chunk({}, finish_reason="tool_calls"),
            ]
        else:
            events = [
                _chunk({"role": "assistant"}),
                _chunk({"content": "Deterministic local fixture complete."}),
                _chunk({}, finish_reason="stop"),
            ]
        self._write(200, "text/event-stream", _sse(events))


class LocalOpenCodeServer(ThreadingHTTPServer):
    expected_model: str
    mode: str
    requests: list[dict]


def create_server(
    host: str,
    port: int,
    *,
    expected_model: str = DEFAULT_REQUEST_MODEL,
    mode: str = "success",
) -> LocalOpenCodeServer:
    if mode not in {"success", "provider-error"}:
        raise ValueError(f"unsupported fixture mode: {mode}")
    server = LocalOpenCodeServer((host, port), LocalOpenCodeHandler)
    server.expected_model = expected_model
    server.mode = mode
    server.requests = []
    return server


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--host", default="0.0.0.0")
    parser.add_argument("--port", type=int, default=0)
    parser.add_argument("--expected-model", default=DEFAULT_REQUEST_MODEL)
    parser.add_argument("--mode", choices=("success", "provider-error"), default="success")
    parser.add_argument("--port-file", type=Path)
    args = parser.parse_args()

    server = create_server(
        args.host,
        args.port,
        expected_model=args.expected_model,
        mode=args.mode,
    )
    if args.port_file:
        args.port_file.write_text(f"{server.server_port}\n")
    print(f"Local OpenCode fixture listening on {args.host}:{server.server_port}", flush=True)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
