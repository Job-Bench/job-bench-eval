"""Exercise every task container without a language model or a real judge.

This fixture validates file parsers and records observations. Its submissions
are diagnostic artifacts, not solutions, and synthetic rewards are not scores.
"""

from __future__ import annotations

from harbor.agents.nop import NopAgent


INSPECTOR = r'''
import csv
import hashlib
import json
import os
from pathlib import Path
import shutil
import sqlite3
import urllib.request
import warnings
from xml.etree import ElementTree
import zipfile

warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')
root = Path('/workspace/task_folder')
output = Path('/workspace/output')
report = {'fixture': True, 'files': [], 'errors': [], 'isolation': {}, 'tools': {}}
assert Path.cwd() == Path('/workspace'), f'Unexpected working directory: {Path.cwd()}'
assert (root / 'TASK_INSTRUCTIONS.txt').is_file()
assert not Path('/tests').exists(), 'Verifier files exposed to agent'
assert not any(name.startswith('JUDGE_') for name in os.environ), 'Judge environment exposed'
for filename in ('RUBRICS.json', 'task_card.md', 'files_required_to_search'):
    assert not list(Path('/workspace').rglob(filename)), f'Hidden material exposed: {filename}'
report['isolation'] = {'workdir': str(Path.cwd()), 'hidden_material_absent': True,
                       'judge_environment_absent': True}
for tool in ('bash', 'curl', 'file', 'git', 'jq', 'sqlite3', 'pdftotext',
             'pdftoppm', 'tesseract', 'libreoffice', 'pandoc', 'ogrinfo'):
    report['tools'][tool] = shutil.which(tool) is not None
    assert report['tools'][tool], f'Missing advertised tool: {tool}'
try:
    with urllib.request.urlopen('https://huggingface.co/robots.txt', timeout=20) as response:
        report['public_network'] = {'ok': response.status == 200, 'status': response.status}
except Exception as exc:
    report['public_network'] = {'ok': False, 'error': f'{type(exc).__name__}: {exc}'}

for path in sorted(root.rglob('*')):
    assert not path.is_symlink(), f'Unexpected symlink: {path}'
    if not path.is_file():
        continue
    data = path.read_bytes()
    suffix = path.suffix.lower()
    item = {'path': path.relative_to(root).as_posix(), 'bytes': len(data),
            'sha256': hashlib.sha256(data).hexdigest(), 'suffix': suffix, 'status': 'ok'}
    try:
        if suffix in ('.xlsx', '.docx'):
            with zipfile.ZipFile(path) as archive:
                assert archive.testzip() is None, 'Office ZIP checksum failed'
        if suffix == '.xlsx':
            import openpyxl
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
            try:
                item['sheets'] = [{'name': s.title, 'rows': s.max_row, 'columns': s.max_column,
                                   'first_row_cells': len(next(s.iter_rows(max_row=1), ()))}
                                  for s in workbook.worksheets]
            finally:
                workbook.close()
        elif suffix == '.xls':
            import xlrd
            workbook = xlrd.open_workbook(path, on_demand=True)
            try:
                item['sheets'] = [{'name': s.name, 'rows': s.nrows, 'columns': s.ncols}
                                  for s in workbook.sheets()]
            finally:
                workbook.release_resources()
        elif suffix == '.docx':
            from docx import Document
            document = Document(path)
            item.update(paragraphs=len(document.paragraphs), tables=len(document.tables))
        elif suffix == '.pdf':
            import fitz
            with fitz.open(path) as document:
                assert not document.needs_pass, 'PDF requires a password'
                item.update(pages=len(document), repaired=document.is_repaired)
                item['page_text_characters'] = [len(page.get_text()) for page in document]
                item['pages_without_text'] = sum(n == 0 for n in item['page_text_characters'])
        elif suffix in ('.png', '.jpg', '.jpeg', '.tif', '.tiff'):
            from PIL import Image
            with Image.open(path) as image:
                image.verify()
            with Image.open(path) as image:
                image.load()
                item['dimensions'] = list(image.size)
        elif suffix == '.db':
            with sqlite3.connect(f'file:{path}?mode=ro', uri=True) as connection:
                connection.execute('PRAGMA query_only = ON')
                check = connection.execute('PRAGMA quick_check').fetchall()
                assert check == [('ok',)], f'SQLite quick_check: {check}'
                item['tables'] = [r[0] for r in connection.execute(
                    "SELECT name FROM sqlite_master WHERE type='table'")]
        elif suffix == '.csv':
            encoding = 'utf-8-sig'
            try:
                data.decode(encoding)
            except UnicodeDecodeError:
                encoding = 'utf-16' if data.startswith((b'\xff\xfe', b'\xfe\xff')) else 'cp1252'
                data.decode(encoding)
            with path.open(encoding=encoding, newline='') as handle:
                rows = list(csv.reader(handle))
            item.update(encoding=encoding, rows=len(rows),
                        column_counts=sorted({len(row) for row in rows}))
        elif suffix in ('.json', '.geojson'):
            parsed = json.loads(data)
            item['json_type'] = type(parsed).__name__
            if suffix == '.geojson':
                import geopandas
                frame = geopandas.read_file(path)
                item.update(features=len(frame), crs=str(frame.crs))
        elif suffix == '.xml':
            item['root_tag'] = ElementTree.fromstring(data).tag
        elif suffix in ('.yaml', '.yml'):
            import yaml
            item['documents'] = len(list(yaml.safe_load_all(data)))
        elif suffix == '.stl':
            import trimesh
            mesh = trimesh.load(path, force='mesh')
            item.update(vertices=len(mesh.vertices), faces=len(mesh.faces))
        elif suffix in ('.txt', '.md', '.rules', '.conf', '.html'):
            item['characters'] = len(data.decode('utf-8-sig'))
        else:
            item['status'] = 'unhandled'
            item['error'] = f'No file parser configured for {suffix}'
    except Exception as exc:
        item.update(status='error', error=f'{type(exc).__name__}: {exc}')
    if item['status'] != 'ok':
        report['errors'].append({'path': item['path'], 'error': item['error']})
    report['files'].append(item)

report['input_files'] = len(report['files'])
encoded = json.dumps(report, ensure_ascii=False, indent=2) + '\n'
logs = Path('/logs/agent/task-audit.json')
logs.write_text(encoded)
logs.chmod(0o644)
output.mkdir(exist_ok=True)
(output / 'harbor-audit.txt').write_text('DIAGNOSTIC FIXTURE, NOT A TASK SOLUTION\n' + encoded)
from PIL import Image
Image.new('RGB', (8, 8), 'blue').save(output / 'harbor-audit.png')
print(json.dumps({'input_files': report['input_files'], 'parse_errors': len(report['errors']),
                  'network': report['public_network']['ok']}))
'''


class AuditAgent(NopAgent):
    @staticmethod
    def name() -> str:
        return 'jobbench-task-audit-fixture'

    async def run(self, instruction, environment, context) -> None:
        result = await environment.exec(
            command="python - <<'PY'\n" + INSPECTOR + '\nPY', timeout_sec=300,
        )
        if result.return_code:
            raise RuntimeError(f'Task audit failed: {result.stdout}\n{result.stderr}')
        context.cost_usd = 0.0
        context.n_input_tokens = 0
        context.n_output_tokens = 0
        context.metadata = {'fixture': True, 'audit_summary': result.stdout.strip()}
